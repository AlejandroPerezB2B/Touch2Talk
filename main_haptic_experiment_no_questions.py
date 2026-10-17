# -*- coding: utf-8 -*-
"""
main_haptic_experiment_no_questions.py

Main experiment for the haptic entrainment study.

Core design
-----------
- Three fixed stimulus lists (List1, List2, List3), 84 trials each.
- Six counterbalancing versions determine whether each block is:
    Lip / Fingertip / No stimulation.
- One continuous four-talker masker.
- Masker starts 1000 ms before target onset and ends 1000 ms after target offset.
- 500-ms raised-cosine masker ramp at onset and offset.
- Participant-specific TMR entered once and kept fixed for the entire experiment.
- Stereo output:
    channel 1 = masker + target sentence at fixed TMR
    channel 2 = 150-Hz speech-derived haptic waveform, or digital zeros
- Haptic and target begin at exactly the same sample, 1000 ms after masker onset.
- One unique EEG event code for every main-experiment sentence (1..252).
- TriggerBox hardware call is isolated in send_eeg_trigger(). It is initially in TEST mode.
- Response beep after masker offset.
- After beep offset, microphone records exactly 3.000 s using the default Windows input.
- One response WAV is saved for every trial.
- ITI = 3000 ms.

Required files
--------------
1. haptic_stimulus_lists.xlsx
   Contains List1, List2, List3, and ConditionOrders sheets.
2. Folder named PsychoPy_target_audio_44100Hz containing target-sentence WAVs.
3. Folder named haptic_wav containing haptic WAVs named, e.g. 5.25_haptic150.wav.
4. One continuous four-talker masker WAV.
5. One response-beep WAV.

Audio files are expected at 44.1 kHz.
"""

from pathlib import Path
from datetime import datetime
import csv
import json
import math
import random
import secrets
import traceback
import tkinter as tk
from tkinter import filedialog

import numpy as np
from scipy.io import wavfile
from openpyxl import load_workbook

from psychopy import prefs
prefs.hardware["audioLib"] = ["ptb"]

from psychopy import core, event, gui, sound, visual

try:
    import sounddevice as sd
except ImportError as exc:
    raise ImportError(
        "The Python package 'sounddevice' is required for microphone recording."
    ) from exc

try:
    import soundfile as sf
    HAVE_SOUNDFILE = True
except ImportError:
    HAVE_SOUNDFILE = False


# ============================================================
# CONSTANTS
# ============================================================

FS = 44100

PRE_MASKER_S = 1.000
POST_MASKER_S = 1.000
MASKER_RAMP_S = 0.500

FIXATION_BEFORE_MASKER_S = 0.500
MIC_RECORDING_S = 3.000
ITI_S = 3.000

N_BLOCKS = 3
TRIALS_PER_BLOCK = 84
TOTAL_TRIALS = N_BLOCKS * TRIALS_PER_BLOCK

# Same auditory reference as the calibration program.
REFERENCE_RMS_DBFS = -30.0
AUDITORY_MASTER_GAIN = 0.70

# EEG trigger mode.
# Keep "test" until the Brain Products TriggerBox interface is known.
# In test mode the trigger is logged/printed but no hardware command is sent.
TRIGGER_MODE = "test"


# ============================================================
# GENERAL UTILITIES
# ============================================================

def sanitize_text_for_filename(text):
    text = str(text).strip()
    out = []
    for ch in text:
        if ch.isalnum() or ch in ("-", "_", "."):
            out.append(ch)
        else:
            out.append("_")
    cleaned = "".join(out).strip("_")
    return cleaned or "unknown"


def db_to_amp(db):
    return 10.0 ** (float(db) / 20.0)


def rms(x):
    x = np.asarray(x, dtype=np.float64)
    if x.size == 0:
        return 0.0
    return float(np.sqrt(np.mean(x * x)))


def to_float_audio(data):
    if np.issubdtype(data.dtype, np.floating):
        return data.astype(np.float64)
    if data.dtype == np.int16:
        return data.astype(np.float64) / 32768.0
    if data.dtype == np.int32:
        # scipy uses int32 container for 24-bit WAVs as well.
        return data.astype(np.float64) / 2147483648.0
    if data.dtype == np.uint8:
        return (data.astype(np.float64) - 128.0) / 128.0
    raise TypeError(f"Unsupported WAV dtype: {data.dtype}")


def make_mono(x):
    x = np.asarray(x, dtype=np.float64)
    if x.ndim == 1:
        return x
    if x.ndim == 2:
        return np.mean(x, axis=1)
    raise ValueError(f"Unexpected audio shape: {x.shape}")


def read_mono_wav(path):
    fs, data = wavfile.read(str(path))
    if fs != FS:
        raise RuntimeError(f"{path.name} is {fs} Hz; expected {FS} Hz.")
    x = make_mono(to_float_audio(data))
    if x.size == 0:
        raise RuntimeError(f"{path.name} is empty.")
    if rms(x) <= 0:
        raise RuntimeError(f"{path.name} is silent.")
    return x


def raised_cosine(n_samples):
    if n_samples <= 1:
        return np.ones(max(1, n_samples), dtype=np.float64)
    return 0.5 - 0.5 * np.cos(np.linspace(0, np.pi, n_samples))


def choose_folder(title):
    root = tk.Tk()
    root.withdraw()
    root.attributes("-topmost", True)
    folder = filedialog.askdirectory(title=title)
    root.destroy()
    return Path(folder) if folder else None


def choose_named_folder(title, expected_name):
    folder = choose_folder(title)
    if folder is None:
        return None
    if folder.name.lower() != expected_name.lower():
        raise RuntimeError(
            f"Incorrect folder selected. Expected a folder named "
            f"'{expected_name}', but selected:\n{folder}"
        )
    return folder


def choose_file(title, filetypes=None):
    root = tk.Tk()
    root.withdraw()
    root.attributes("-topmost", True)
    filename = filedialog.askopenfilename(
        title=title,
        filetypes=filetypes or [("All files", "*.*")],
    )
    root.destroy()
    return Path(filename) if filename else None


# ============================================================
# STARTUP
# ============================================================

def startup_dialog():
    dlg = gui.Dlg(title="Experimento principal")
    dlg.addField("Participant ID:")
    dlg.addField("Condition-order version (1-6):", choices=[1, 2, 3, 4, 5, 6])
    dlg.addField("Participant-specific TMR (dB, keep sign):", initial="0.0")
    values = dlg.show()

    if not dlg.OK or values is None:
        return None

    participant_id = sanitize_text_for_filename(values[0])
    version = int(values[1])

    try:
        tmr_db = float(values[2])
    except Exception as exc:
        raise ValueError("TMR must be a signed numeric value, e.g. -3.5 or +2.0.") from exc

    list_workbook = choose_file(
        "Select haptic_stimulus_lists.xlsx",
        [("Excel workbook", "*.xlsx")],
    )
    if list_workbook is None:
        return None

    target_folder = choose_named_folder(
        "Select the folder named PsychoPy_target_audio_44100Hz",
        "PsychoPy_target_audio_44100Hz",
    )
    if target_folder is None:
        return None

    haptic_folder = choose_named_folder(
        "Select the folder named haptic_wav",
        "haptic_wav",
    )
    if haptic_folder is None:
        return None

    masker_file = choose_file(
        "Select continuous four-talker masker WAV",
        [("WAV audio", "*.wav")],
    )
    if masker_file is None:
        return None

    beep_file = choose_file(
        "Select response beep WAV",
        [("WAV audio", "*.wav")],
    )
    if beep_file is None:
        return None

    output_folder = choose_folder(
        "Select folder where experiment results will be saved"
    )
    if output_folder is None:
        return None

    return {
        "participant_id": participant_id,
        "version": version,
        "tmr_db": tmr_db,
        "list_workbook": list_workbook,
        "target_folder": target_folder,
        "haptic_folder": haptic_folder,
        "masker_file": masker_file,
        "beep_file": beep_file,
        "output_folder": output_folder,
    }


# ============================================================
# EXCEL INPUT
# ============================================================

def read_stimulus_workbook(path, version):
    wb = load_workbook(path, data_only=True, read_only=True)

    required_sheets = ["List1", "List2", "List3", "ConditionOrders"]
    missing_sheets = [s for s in required_sheets if s not in wb.sheetnames]
    if missing_sheets:
        raise RuntimeError(
            "Missing sheet(s) in stimulus workbook: " + ", ".join(missing_sheets)
        )

    lists = {}

    for list_index, sheet_name in enumerate(("List1", "List2", "List3"), start=1):
        ws = wb[sheet_name]
        headers = [cell.value for cell in next(ws.iter_rows(min_row=1, max_row=1))]
        header_map = {str(h).strip(): i for i, h in enumerate(headers) if h is not None}

        required_headers = [
            "Trial", "Code", "Sentence", "Predictability", "TargetFamily",
            "SentenceNumber", "SimplifiedFilename"
        ]
        missing_headers = [h for h in required_headers if h not in header_map]
        if missing_headers:
            raise RuntimeError(
                f"{sheet_name} is missing column(s): {', '.join(missing_headers)}"
            )

        rows = []
        for row in ws.iter_rows(min_row=2, values_only=True):
            if row[header_map["Trial"]] is None:
                continue

            trial = int(row[header_map["Trial"]])
            if trial < 1 or trial > TRIALS_PER_BLOCK:
                raise RuntimeError(f"Unexpected trial number {trial} in {sheet_name}.")

            trigger_code = (list_index - 1) * TRIALS_PER_BLOCK + trial

            rows.append({
                "list_name": sheet_name,
                "block": list_index,
                "trial": trial,
                "global_trial": (list_index - 1) * TRIALS_PER_BLOCK + trial,
                "trigger_code": trigger_code,
                "code": str(row[header_map["Code"]]).strip(),
                "sentence": str(row[header_map["Sentence"]]).strip(),
                "predictability": str(row[header_map["Predictability"]]).strip(),
                "target_family": row[header_map["TargetFamily"]],
                "sentence_number": row[header_map["SentenceNumber"]],
                "target_file": str(row[header_map["SimplifiedFilename"]]).strip(),
            })

        rows.sort(key=lambda x: x["trial"])

        if len(rows) != TRIALS_PER_BLOCK:
            raise RuntimeError(
                f"{sheet_name} contains {len(rows)} usable trials; expected {TRIALS_PER_BLOCK}."
            )

        expected_trials = list(range(1, TRIALS_PER_BLOCK + 1))
        observed_trials = [r["trial"] for r in rows]
        if observed_trials != expected_trials:
            raise RuntimeError(f"Trial order in {sheet_name} is not exactly 1..84.")

        lists[sheet_name] = rows

    # Read counterbalancing table directly from workbook.
    ws = wb["ConditionOrders"]
    headers = [cell.value for cell in next(ws.iter_rows(min_row=1, max_row=1))]
    header_map = {str(h).strip(): i for i, h in enumerate(headers) if h is not None}

    needed = ["Version", "List1 condition", "List2 condition", "List3 condition"]
    missing = [h for h in needed if h not in header_map]
    if missing:
        raise RuntimeError(
            "ConditionOrders is missing column(s): " + ", ".join(missing)
        )

    condition_map = None
    for row in ws.iter_rows(min_row=2, values_only=True):
        if row[header_map["Version"]] is None:
            continue
        if int(row[header_map["Version"]]) == int(version):
            condition_map = {
                1: str(row[header_map["List1 condition"]]).strip(),
                2: str(row[header_map["List2 condition"]]).strip(),
                3: str(row[header_map["List3 condition"]]).strip(),
            }
            break

    if condition_map is None:
        raise RuntimeError(f"Version {version} was not found in ConditionOrders.")

    valid_conditions = {"Lip", "Fingertip", "No stimulation"}
    for block, condition in condition_map.items():
        if condition not in valid_conditions:
            raise RuntimeError(
                f"Unexpected condition '{condition}' for block {block}."
            )

    return lists, condition_map



# ============================================================
# PRELOAD AUDIO
# ============================================================

def preload_audio(lists, target_folder, haptic_folder, masker_file, beep_file):
    target_audio = {}
    haptic_audio = {}

    all_trials = lists["List1"] + lists["List2"] + lists["List3"]

    for item in all_trials:
        target_file = item["target_file"]
        target_path = target_folder / target_file

        if not target_path.exists():
            raise RuntimeError(f"Missing target WAV: {target_path}")

        if target_file not in target_audio:
            target_audio[target_file] = read_mono_wav(target_path)

        stem = Path(target_file).stem
        haptic_file = f"{stem}_haptic150.wav"
        haptic_path = haptic_folder / haptic_file

        if not haptic_path.exists():
            raise RuntimeError(f"Missing haptic WAV: {haptic_path}")

        if haptic_file not in haptic_audio:
            haptic_audio[haptic_file] = read_mono_wav(haptic_path)

    long_masker = read_mono_wav(masker_file)
    beep = read_mono_wav(beep_file)

    return target_audio, haptic_audio, long_masker, beep


# ============================================================
# AUDIO CONSTRUCTION
# ============================================================

def build_stereo_trial(target, haptic, long_masker, tmr_db, haptic_present, rng):
    pre_n = int(round(PRE_MASKER_S * FS))
    post_n = int(round(POST_MASKER_S * FS))
    ramp_n = int(round(MASKER_RAMP_S * FS))

    target_n = len(target)
    total_n = pre_n + target_n + post_n

    if len(long_masker) < total_n:
        raise RuntimeError(
            "Continuous masker is shorter than the required trial duration."
        )

    max_start = len(long_masker) - total_n
    masker_start = rng.randint(0, max_start) if max_start > 0 else 0
    masker = long_masker[masker_start:masker_start + total_n].copy()

    reference_rms = db_to_amp(REFERENCE_RMS_DBFS)

    # Define fixed masker level using the region simultaneous with the target.
    masker_overlap = masker[pre_n:pre_n + target_n]
    masker_overlap_rms = rms(masker_overlap)
    if masker_overlap_rms <= 0:
        raise RuntimeError("Selected masker segment has zero RMS.")
    masker *= reference_rms / masker_overlap_rms

    # Apply 500-ms onset/offset ramps inside the 1-s pre/post periods.
    ramp = raised_cosine(ramp_n)
    masker[:ramp_n] *= ramp
    masker[-ramp_n:] *= ramp[::-1]

    # Normalize target to the same RMS reference, then apply signed TMR.
    target_rms = rms(target)
    if target_rms <= 0:
        raise RuntimeError("Target sentence has zero RMS.")

    target_norm = target * (reference_rms / target_rms)
    target_gain = db_to_amp(tmr_db)
    target_scaled = target_norm * target_gain

    target_padded = np.zeros(total_n, dtype=np.float64)
    target_padded[pre_n:pre_n + target_n] = target_scaled

    auditory = (masker + target_padded) * AUDITORY_MASTER_GAIN

    auditory_peak = float(np.max(np.abs(auditory)))
    if auditory_peak >= 1.0:
        raise RuntimeError(
            f"Auditory channel clipping risk: peak={auditory_peak:.3f}."
        )

    # Haptic file must correspond to target timing. If file duration differs,
    # crop/pad only at the end; onset remains sample-identical to target onset.
    haptic_target = np.asarray(haptic, dtype=np.float64)
    if len(haptic_target) > target_n:
        haptic_target = haptic_target[:target_n]
    elif len(haptic_target) < target_n:
        haptic_target = np.pad(haptic_target, (0, target_n - len(haptic_target)))

    haptic_channel = np.zeros(total_n, dtype=np.float64)
    if haptic_present:
        haptic_channel[pre_n:pre_n + target_n] = haptic_target

    haptic_peak = float(np.max(np.abs(haptic_channel))) if haptic_channel.size else 0.0
    if haptic_peak >= 1.0:
        raise RuntimeError(
            f"Haptic channel clipping risk: peak={haptic_peak:.3f}."
        )

    stereo = np.column_stack((auditory, haptic_channel))

    meta = {
        "masker_start_sample": int(masker_start),
        "masker_start_s": float(masker_start / FS),
        "target_onset_sample": int(pre_n),
        "target_onset_s": PRE_MASKER_S,
        "target_duration_s": float(target_n / FS),
        "trial_audio_duration_s": float(total_n / FS),
        "target_gain_linear": float(target_gain),
        "auditory_peak": auditory_peak,
        "haptic_peak": haptic_peak,
        "haptic_present": int(bool(haptic_present)),
    }

    return stereo, meta


def build_beep_stereo(beep):
    # Beep is sent only to auditory channel 1. Channel 2 remains zero so the
    # response cue is never sent to the tactor.
    return np.column_stack((beep, np.zeros(len(beep), dtype=np.float64)))


# ============================================================
# EEG TRIGGERS
# ============================================================

def send_eeg_trigger(trigger_code):
    """
    Send the target-onset marker.

    CURRENT STATE: TEST MODE ONLY.

    Replace the body of the hardware branch when the Brain Products
    TriggerBox connection/API is known. The rest of the experiment does not
    need to change.
    """
    trigger_code = int(trigger_code)

    if not (1 <= trigger_code <= 252):
        raise ValueError(f"Invalid target trigger code: {trigger_code}")

    if TRIGGER_MODE == "test":
        print(f"[EEG TEST] trigger {trigger_code}")
        return

    if TRIGGER_MODE == "hardware":
        raise NotImplementedError(
            "TriggerBox hardware communication has not yet been configured."
        )

    raise ValueError(f"Unknown TRIGGER_MODE: {TRIGGER_MODE}")


# ============================================================
# MICROPHONE
# ============================================================

def record_microphone(output_path, duration_s=MIC_RECORDING_S):
    n_samples = int(round(duration_s * FS))

    recording = sd.rec(
        n_samples,
        samplerate=FS,
        channels=1,
        dtype="float32",
        blocking=False,
    )
    sd.wait()

    recording = np.asarray(recording[:, 0], dtype=np.float32)

    if HAVE_SOUNDFILE:
        sf.write(str(output_path), recording, FS, subtype="PCM_24")
    else:
        # Fallback: valid 32-bit float WAV.
        wavfile.write(str(output_path), FS, recording)

    return {
        "response_recording_samples": int(len(recording)),
        "response_recording_duration_s": float(len(recording) / FS),
        "response_recording_peak": float(np.max(np.abs(recording))) if len(recording) else 0.0,
    }


# ============================================================
# DISPLAY
# ============================================================

def draw_text(win, text, color="white", height=0.045):
    stim = visual.TextStim(
        win,
        text=text,
        color=color,
        height=height,
        wrapWidth=1.65,
        alignText="center",
    )
    stim.draw()
    win.flip()
    return stim


def wait_for_space(win, text, color="white"):
    draw_text(win, text, color=color)
    key = event.waitKeys(keyList=["space", "escape"])[0]
    if key == "escape":
        raise KeyboardInterrupt


def show_fixation(win):
    fixation = visual.TextStim(win, text="+", color="white", height=0.08)
    fixation.draw()
    win.flip()


def show_block_screen(win, block, condition):
    condition_spanish = {
        "Lip": "LABIO SUPERIOR",
        "Fingertip": "YEMAS DE LOS DEDOS",
        "No stimulation": "SIN ESTIMULACIÓN",
    }[condition]

    wait_for_space(
        win,
        (
            f"BLOQUE {block} DE 3\n\n"
            f"CONDICIÓN: {condition_spanish}\n\n"
            "Descanse el tiempo que necesite.\n\n"
            "Avise al experimentador cuando esté preparado para continuar.\n\n"
            "Pulse ESPACIO para continuar."
        ),
    )


def show_response_screen(win):
    draw_text(
        win,
        "REPITA EN VOZ ALTA LA ÚLTIMA PALABRA",
        color="white",
        height=0.05,
    )



def run_iti(win):
    show_fixation(win)
    core.wait(ITI_S)


# ============================================================
# DATA SAVING
# ============================================================

TRIAL_FIELDS = [
    "participant_id",
    "version",
    "tmr_db",
    "block",
    "list",
    "condition",
    "trial_in_block",
    "global_trial",
    "trigger_code",
    "code",
    "target_file",
    "haptic_file",
    "predictability",
    "target_family",
    "sentence_number",
    "haptic_present",
    "masker_start_sample",
    "masker_start_s",
    "target_onset_sample",
    "target_onset_s",
    "target_duration_s",
    "trial_audio_duration_s",
    "target_gain_linear",
    "auditory_peak",
    "haptic_peak",
    "trigger_mode",
    "trigger_perf_time_s",
    "response_file",
    "response_recording_samples",
    "response_recording_duration_s",
    "response_recording_peak",
    "timestamp",
]


def append_trial(csv_path, row):
    exists = csv_path.exists()
    with csv_path.open("a", newline="", encoding="utf-8-sig") as f:
        writer = csv.DictWriter(f, fieldnames=TRIAL_FIELDS)
        if not exists:
            writer.writeheader()
        writer.writerow(row)


def write_trigger_mapping(path, lists):
    fields = [
        "trigger_code", "block", "list", "trial_in_block", "code", "target_file",
        "predictability", "target_family", "sentence_number"
    ]
    with path.open("w", newline="", encoding="utf-8-sig") as f:
        writer = csv.DictWriter(f, fieldnames=fields)
        writer.writeheader()
        for list_name in ("List1", "List2", "List3"):
            for item in lists[list_name]:
                writer.writerow({
                    "trigger_code": item["trigger_code"],
                    "block": item["block"],
                    "list": item["list_name"],
                    "trial_in_block": item["trial"],
                    "code": item["code"],
                    "target_file": item["target_file"],
                    "predictability": item["predictability"],
                    "target_family": item["target_family"],
                    "sentence_number": item["sentence_number"],
                })


# ============================================================
# MAIN EXPERIMENT
# ============================================================

def run_experiment():
    info = startup_dialog()
    if info is None:
        print("Experiment cancelled before start.")
        return

    participant_id = info["participant_id"]
    version = info["version"]
    tmr_db = info["tmr_db"]
    output_root = info["output_folder"]

    stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    session_dir = output_root / f"{participant_id}_V{version:02d}_{stamp}"
    response_dir = session_dir / "spoken_responses"
    session_dir.mkdir(parents=True, exist_ok=True)
    response_dir.mkdir(parents=True, exist_ok=True)

    print("Loading experiment lists...")
    lists, condition_map = read_stimulus_workbook(
        info["list_workbook"], version
    )

    print("Preloading target, haptic, masker, and beep audio...")
    target_audio, haptic_audio, long_masker, beep = preload_audio(
        lists=lists,
        target_folder=info["target_folder"],
        haptic_folder=info["haptic_folder"],
        masker_file=info["masker_file"],
        beep_file=info["beep_file"],
    )

    beep_stereo = build_beep_stereo(beep)
    beep_sound = sound.Sound(
        value=beep_stereo,
        sampleRate=FS,
        stereo=True,
        hamming=False,
    )
    beep_duration_s = len(beep) / FS

    seed = secrets.randbits(32)
    rng = random.Random(seed)

    trials_csv = session_dir / "trial_log.csv"
    summary_json = session_dir / "session_summary.json"
    trigger_mapping_csv = session_dir / "trigger_mapping.csv"
    write_trigger_mapping(trigger_mapping_csv, lists)

    summary = {
        "participant_id": participant_id,
        "version": version,
        "tmr_db": tmr_db,
        "condition_by_block": condition_map,
        "random_seed": seed,
        "trigger_mode": TRIGGER_MODE,
        "trigger_mapping": "List1=1..84, List2=85..168, List3=169..252",
        "stimulus_workbook": str(info["list_workbook"]),
        "target_folder": str(info["target_folder"]),
        "haptic_folder": str(info["haptic_folder"]),
        "masker_file": str(info["masker_file"]),
        "beep_file": str(info["beep_file"]),
        "audio_sample_rate_hz": FS,
        "pre_masker_s": PRE_MASKER_S,
        "post_masker_s": POST_MASKER_S,
        "masker_ramp_s": MASKER_RAMP_S,
        "fixation_before_masker_s": FIXATION_BEFORE_MASKER_S,
        "microphone_recording_s": MIC_RECORDING_S,
        "iti_s": ITI_S,
        "comprehension_questions": False,
        "started": datetime.now().isoformat(timespec="seconds"),
        "completed": None,
        "aborted": False,
    }

    with summary_json.open("w", encoding="utf-8") as f:
        json.dump(summary, f, indent=2, ensure_ascii=False)

    win = visual.Window(
        fullscr=True,
        color="black",
        units="height",
        allowGUI=False,
    )

    try:
        wait_for_space(
            win,
            (
                "EXPERIMENTO PRINCIPAL\n\n"
                "Escuchará frases con ruido de fondo.\n\n"
                "Cuando oiga el tono, repita en voz alta la última palabra, "
                "incluso si no está seguro de la respuesta.\n\n"
                "Pulse ESPACIO para continuar."
            ),
        )

        for block in range(1, N_BLOCKS + 1):
            list_name = f"List{block}"
            condition = condition_map[block]
            haptic_present = condition != "No stimulation"

            show_block_screen(win, block, condition)

            for item in lists[list_name]:
                trial = item["trial"]
                global_trial = item["global_trial"]
                trigger_code = item["trigger_code"]
                target_file = item["target_file"]
                stem = Path(target_file).stem
                haptic_file = f"{stem}_haptic150.wav"

                print(
                    f"Block {block}/3 | Trial {trial}/84 | "
                    f"{target_file} | {condition} | trigger {trigger_code}"
                )

                target = target_audio[target_file]
                haptic = haptic_audio[haptic_file]

                stereo, audio_meta = build_stereo_trial(
                    target=target,
                    haptic=haptic,
                    long_masker=long_masker,
                    tmr_db=tmr_db,
                    haptic_present=haptic_present,
                    rng=rng,
                )

                trial_sound = sound.Sound(
                    value=stereo,
                    sampleRate=FS,
                    stereo=True,
                    hamming=False,
                )

                # Fixation appears 500 ms before masker onset and stays visible
                # throughout the complete auditory stimulus.
                show_fixation(win)
                core.wait(FIXATION_BEFORE_MASKER_S)

                event.clearEvents()
                play_clock = core.Clock()
                trial_sound.play()

                # The target and haptic start 1000 ms after masker onset.
                # Trigger hardware implementation can later replace the isolated
                # function above without changing trial logic.
                core.wait(PRE_MASKER_S)
                send_eeg_trigger(trigger_code)
                trigger_perf_time_s = play_clock.getTime()

                remaining_audio_s = max(
                    0.0,
                    audio_meta["trial_audio_duration_s"] - PRE_MASKER_S,
                )
                core.wait(remaining_audio_s)
                trial_sound.stop()

                # Response cue beep: auditory channel only.
                beep_sound.play()
                core.wait(beep_duration_s)
                beep_sound.stop()

                # Microphone opens only after beep offset.
                show_response_screen(win)

                response_filename = (
                    f"{participant_id}_V{version:02d}_B{block:02d}_"
                    f"T{trial:03d}_{sanitize_text_for_filename(stem)}_response.wav"
                )
                response_path = response_dir / response_filename

                mic_meta = record_microphone(
                    response_path,
                    duration_s=MIC_RECORDING_S,
                )

                row = {
                    "participant_id": participant_id,
                    "version": version,
                    "tmr_db": tmr_db,
                    "block": block,
                    "list": list_name,
                    "condition": condition,
                    "trial_in_block": trial,
                    "global_trial": global_trial,
                    "trigger_code": trigger_code,
                    "code": item["code"],
                    "target_file": target_file,
                    "haptic_file": haptic_file,
                    "predictability": item["predictability"],
                    "target_family": item["target_family"],
                    "sentence_number": item["sentence_number"],
                    "haptic_present": audio_meta["haptic_present"],
                    "masker_start_sample": audio_meta["masker_start_sample"],
                    "masker_start_s": audio_meta["masker_start_s"],
                    "target_onset_sample": audio_meta["target_onset_sample"],
                    "target_onset_s": audio_meta["target_onset_s"],
                    "target_duration_s": audio_meta["target_duration_s"],
                    "trial_audio_duration_s": audio_meta["trial_audio_duration_s"],
                    "target_gain_linear": audio_meta["target_gain_linear"],
                    "auditory_peak": audio_meta["auditory_peak"],
                    "haptic_peak": audio_meta["haptic_peak"],
                    "trigger_mode": TRIGGER_MODE,
                    "trigger_perf_time_s": trigger_perf_time_s,
                    "response_file": response_filename,
                    "response_recording_samples": mic_meta["response_recording_samples"],
                    "response_recording_duration_s": mic_meta["response_recording_duration_s"],
                    "response_recording_peak": mic_meta["response_recording_peak"],
                    "timestamp": datetime.now().isoformat(timespec="seconds"),
                }

                append_trial(trials_csv, row)

                # Fixed 3000-ms inter-trial interval.
                run_iti(win)

        summary["completed"] = datetime.now().isoformat(timespec="seconds")
        with summary_json.open("w", encoding="utf-8") as f:
            json.dump(summary, f, indent=2, ensure_ascii=False)

        wait_for_space(
            win,
            (
                "EXPERIMENTO COMPLETADO\n\n"
                "Gracias por su participación.\n\n"
                "Pulse ESPACIO para finalizar."
            ),
        )

    except KeyboardInterrupt:
        summary["aborted"] = True
        summary["completed"] = datetime.now().isoformat(timespec="seconds")
        with summary_json.open("w", encoding="utf-8") as f:
            json.dump(summary, f, indent=2, ensure_ascii=False)
        print("Experiment aborted by user.")

    except Exception:
        traceback.print_exc()
        summary["aborted"] = True
        summary["completed"] = datetime.now().isoformat(timespec="seconds")
        with summary_json.open("w", encoding="utf-8") as f:
            json.dump(summary, f, indent=2, ensure_ascii=False)

        try:
            wait_for_space(
                win,
                (
                    "ERROR DURANTE EL EXPERIMENTO\n\n"
                    "Consulte la consola de PsychoPy para ver el error.\n\n"
                    "Pulse ESPACIO para cerrar."
                ),
                color="red",
            )
        except Exception:
            pass
        raise

    finally:
        win.close()


if __name__ == "__main__":
    run_experiment()
